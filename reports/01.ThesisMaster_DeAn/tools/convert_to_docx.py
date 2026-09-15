# -*- coding: utf-8 -*-
"""Build the master's project report DOCX from the merged Markdown source.

The builder follows the academic layout of the retained sample report:
A4 paper, Times New Roman, justified 13 pt body text, numbered headings,
top-centered page numbers, restrained tables, and inline figures.
"""

from __future__ import annotations

import re
import sys
import textwrap
from pathlib import Path
from urllib.parse import unquote

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

try:
    from PIL import Image
except ImportError:  # pragma: no cover - python-docx can still size images
    Image = None


if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


REPORT_DIR = Path(__file__).resolve().parent.parent
OUTPUTS_DIR = REPORT_DIR / "outputs"
INPUT_FILE = OUTPUTS_DIR / "BaoCaoDeAn_ThacSi_NguyenTienDat_Full.md"
OUTPUT_FILE = OUTPUTS_DIR / "BaoCaoDeAn_ThacSi_NguyenTienDat_Full.docx"
LOGO_FILE = REPORT_DIR / "references" / "logo_ptit.png"

FONT_BODY = "Times New Roman"
FONT_CODE = FONT_BODY
PAGE_WIDTH_CM = 21.0
PAGE_HEIGHT_CM = 29.7
TOP_MARGIN_CM = 3.5
BOTTOM_MARGIN_CM = 3.0
LEFT_MARGIN_CM = 3.5
RIGHT_MARGIN_CM = 2.0
HEADER_DISTANCE_CM = 1.5
FOOTER_DISTANCE_CM = 1.5
BODY_SIZE = 13
TABLE_SIZE = BODY_SIZE
TABLE_INDENT_DXA = 120
USABLE_DXA = round((PAGE_WIDTH_CM - LEFT_MARGIN_CM - RIGHT_MARGIN_CM) / 2.54 * 1440)
TABLE_DXA = USABLE_DXA - (2 * TABLE_INDENT_DXA)
MAIN_START_TITLE = "MỞ ĐẦU"

FRONT_TITLES = {"LỜI CẢM ƠN", "LỜI NÓI ĐẦU", "MỤC LỤC"}
TOC_FRONT_TITLES = {
    "DANH MỤC BẢNG",
    "DANH MỤC HÌNH",
    "DANH MỤC CHỮ CÁI VIẾT TẮT",
}


def set_run_font(run, name=FONT_BODY, size=None, bold=None, italic=None):
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "vi-VN")
    lang.set(qn("w:eastAsia"), "vi-VN")


def get_or_add_style(doc, name, style_type=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, style_type)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_BODY
    normal.font.size = Pt(BODY_SIZE)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_BODY)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_BODY)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(0.75)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    # Word creates TOC entries as hyperlinks.  Keep them visually consistent
    # with the monochrome academic sample while retaining click navigation.
    for name in ("Hyperlink", "FollowedHyperlink"):
        style = get_or_add_style(doc, name, WD_STYLE_TYPE.CHARACTER)
        style.font.name = FONT_BODY
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_BODY)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_BODY)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)

    heading_tokens = {
        "Heading 1": (13, True, False, WD_ALIGN_PARAGRAPH.CENTER, 0, 12),
        "Heading 2": (13, True, False, WD_ALIGN_PARAGRAPH.LEFT, 10, 5),
        "Heading 3": (13, True, True, WD_ALIGN_PARAGRAPH.LEFT, 8, 4),
        "Heading 4": (13, True, False, WD_ALIGN_PARAGRAPH.LEFT, 6, 3),
    }
    for name, (size, bold, italic, align, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = FONT_BODY
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_BODY)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_BODY)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        pf = style.paragraph_format
        pf.alignment = align
        pf.first_line_indent = Cm(0)
        pf.left_indent = Cm(0)
        pf.line_spacing = 1.5
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.keep_with_next = True
        pf.keep_together = True

    front = get_or_add_style(doc, "Front Matter Title")
    front.base_style = normal
    front.font.name = FONT_BODY
    front.font.size = Pt(BODY_SIZE)
    front.font.bold = True
    front.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    front.paragraph_format.first_line_indent = Cm(0)
    front.paragraph_format.line_spacing = 1.5
    front.paragraph_format.space_after = Pt(0)
    front.paragraph_format.keep_with_next = True

    part = get_or_add_style(doc, "Part Title")
    part.base_style = normal
    part.font.name = FONT_BODY
    part.font.size = Pt(BODY_SIZE)
    part.font.bold = True
    part.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    part.paragraph_format.first_line_indent = Cm(0)
    part.paragraph_format.line_spacing = 1.5
    part.paragraph_format.space_after = Pt(5)
    part.paragraph_format.keep_with_next = True
    part_ppr = part._element.get_or_add_pPr()
    outline = part_ppr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        part_ppr.append(outline)
    outline.set(qn("w:val"), "0")

    for name in ("Table Caption", "Figure Caption"):
        style = get_or_add_style(doc, name)
        style.base_style = normal
        style.font.name = FONT_BODY
        style.font.size = Pt(BODY_SIZE)
        style.font.bold = True
        style.font.italic = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(5)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = name == "Table Caption"

    code = get_or_add_style(doc, "Code Block")
    code.base_style = normal
    code.font.name = FONT_CODE
    code.font.size = Pt(BODY_SIZE)
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    code.paragraph_format.first_line_indent = Cm(0)
    code.paragraph_format.left_indent = Cm(0.6)
    code.paragraph_format.right_indent = Cm(0.4)
    code.paragraph_format.line_spacing = 1.5
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(3)

    ref = get_or_add_style(doc, "Reference")
    ref.base_style = normal
    ref.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ref.paragraph_format.first_line_indent = Cm(-0.8)
    ref.paragraph_format.left_indent = Cm(0.8)
    ref.paragraph_format.line_spacing = 1.5
    ref.paragraph_format.space_after = Pt(0)

    note = get_or_add_style(doc, "Note")
    note.base_style = normal
    note.font.italic = True
    note.paragraph_format.first_line_indent = Cm(0)
    note.paragraph_format.left_indent = Cm(0.7)
    note.paragraph_format.right_indent = Cm(0.5)

    equation = get_or_add_style(doc, "Equation")
    equation.base_style = normal
    equation.font.name = FONT_BODY
    equation.font.size = Pt(BODY_SIZE)
    equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.first_line_indent = Cm(0)
    equation.paragraph_format.line_spacing = 1.5
    equation.paragraph_format.space_before = Pt(5)
    equation.paragraph_format.space_after = Pt(5)
    equation.paragraph_format.keep_together = True

    for list_name in ("List Bullet", "List Number"):
        style = doc.styles[list_name]
        style.font.name = FONT_BODY
        style.font.size = Pt(BODY_SIZE)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_BODY)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_BODY)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(0)

    manual_number = get_or_add_style(doc, "Manual Numbered List")
    manual_number.base_style = normal
    manual_number.font.name = FONT_BODY
    manual_number.font.size = Pt(BODY_SIZE)
    manual_number._element.rPr.rFonts.set(qn("w:ascii"), FONT_BODY)
    manual_number._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_BODY)
    manual_number._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    manual_number.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    manual_number.paragraph_format.left_indent = Cm(0.75)
    manual_number.paragraph_format.first_line_indent = Cm(-0.75)
    manual_number.paragraph_format.line_spacing = 1.5
    manual_number.paragraph_format.space_after = Pt(0)

    # Word materializes dynamic indexes using these built-in styles.  Define
    # their typography up front so field refreshes cannot reintroduce a
    # smaller font or single-spaced entries.
    for name in ("TOC 1", "TOC 2", "TOC 3", "TOC 4", "Table of Figures"):
        style = get_or_add_style(doc, name)
        style.font.name = FONT_BODY
        style.font.size = Pt(BODY_SIZE)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_BODY)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_BODY)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        style.paragraph_format.line_spacing = 1.5


def configure_section(section, cover=False):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(PAGE_WIDTH_CM)
    section.page_height = Cm(PAGE_HEIGHT_CM)
    section.top_margin = Cm(TOP_MARGIN_CM)
    section.bottom_margin = Cm(BOTTOM_MARGIN_CM)
    section.left_margin = Cm(LEFT_MARGIN_CM)
    section.right_margin = Cm(RIGHT_MARGIN_CM)
    section.header_distance = Cm(HEADER_DISTANCE_CM)
    section.footer_distance = Cm(FOOTER_DISTANCE_CM)


def remove_element(parent, tag):
    for element in list(parent.findall(qn(tag))):
        parent.remove(element)


def set_page_border(section, enabled=True):
    sect_pr = section._sectPr
    remove_element(sect_pr, "w:pgBorders")
    if not enabled:
        return
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "double")
        edge.set(qn("w:sz"), "12")
        edge.set(qn("w:space"), "14")
        edge.set(qn("w:color"), "000000")
        borders.append(edge)
    sect_pr.append(borders)


def add_field(paragraph, instruction, display=""):
    run = paragraph.add_run()
    set_run_font(run, size=BODY_SIZE)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = display
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def set_update_fields(doc):
    settings = doc.settings._element
    remove_element(settings, "w:updateFields")
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def clear_story(story):
    """Remove visible header/footer content while retaining one valid paragraph."""
    for table in list(story.tables):
        story._element.remove(table._element)
    if not story.paragraphs:
        story.add_paragraph()
    for paragraph in story.paragraphs:
        ppr = paragraph._p.pPr
        for child in list(paragraph._p):
            if child is not ppr:
                paragraph._p.remove(child)


def prepare_unnumbered_section(section):
    # An unnumbered section should not merely contain a blank story: it must
    # have no header/footer references at all.  Avoid touching section.header
    # or section.footer here because python-docx would create empty parts.
    sect_pr = section._sectPr
    for tag in ("w:headerReference", "w:footerReference", "w:titlePg"):
        remove_element(sect_pr, tag)


def add_page_number(section, start=1):
    prepare_unnumbered_section(section)
    section.header.is_linked_to_previous = False
    header = section.header
    clear_story(header)
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.5
    add_field(p, " PAGE ", str(start))
    sect_pr = section._sectPr
    remove_element(sect_pr, "w:pgNumType")
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:start"), str(start))
    sect_pr.append(pg_num)


def extract_cover_lines(lines):
    values = []
    in_cover = False
    for raw in lines:
        line = raw.strip()
        if line == "<!-- COVER -->":
            in_cover = True
            continue
        if line == "<!-- END_COVER -->":
            break
        if in_cover and line and not line.startswith("!["):
            values.append(strip_inline(line))
    if len(values) != 10:
        raise ValueError(f"Khối COVER phải có đúng 10 dòng dữ liệu, hiện có {len(values)}")
    return values


def add_cover(doc, cover):
    section = doc.sections[0]
    configure_section(section, cover=True)
    prepare_unnumbered_section(section)
    set_page_border(section, True)

    def cover_p(text="", size=BODY_SIZE, bold=False, italic=False, before=0, after=0, spacing=1.5):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = spacing
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic)
        return p

    cover_p(cover[0], BODY_SIZE, True, after=2)
    cover_p("------------------------------------", BODY_SIZE, True, after=6)
    if not LOGO_FILE.exists():
        raise FileNotFoundError(f"Không tìm thấy logo bìa: {LOGO_FILE}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run()
    logo = run.add_picture(str(LOGO_FILE), width=Cm(2.9))
    logo._inline.docPr.set("descr", "Biểu trưng Học viện Công nghệ Bưu chính Viễn thông")
    logo._inline.docPr.set("title", "Biểu trưng PTIT")
    cover_p(cover[1], BODY_SIZE, True, after=16)
    wrapped_title = "\n".join(textwrap.wrap(cover[2], width=58, break_long_words=False))
    cover_p(
        wrapped_title,
        BODY_SIZE,
        True,
        after=4,
        spacing=1.5,
    )
    cover_p(cover[3], BODY_SIZE, True, italic=True, after=16)
    cover_p(cover[4], BODY_SIZE, after=2)
    cover_p(cover[5], BODY_SIZE, after=20)
    cover_p(cover[6], BODY_SIZE, True)
    cover_p(cover[7], BODY_SIZE, True, after=22)
    cover_p(cover[8], BODY_SIZE, True, after=48)
    cover_p(cover[9], BODY_SIZE)

    front_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(front_section, cover=False)
    prepare_unnumbered_section(front_section)
    set_page_border(front_section, False)
    return front_section


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    rpr.append(color)
    rfonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), FONT_BODY)
    rpr.append(rfonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(BODY_SIZE * 2))
    rpr.append(size)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), str(BODY_SIZE * 2))
    rpr.append(size_cs)
    run.append(rpr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text, base_size=BODY_SIZE):
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=base_size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name=FONT_CODE, size=base_size)
        elif token.startswith("["):
            link = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link:
                add_hyperlink(paragraph, link.group(1), link.group(2))
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=base_size, italic=True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=base_size)


def add_body_paragraph(doc, text, style=None, indent=True):
    p = doc.add_paragraph(style=style or "Normal")
    if not indent:
        p.paragraph_format.first_line_indent = Cm(0)
    if style == "Reference":
        # Apply this directly as well as through the style.  Word may retain
        # paragraph-level justification from a refreshed hyperlink/field run,
        # which creates distracting gaps around long DOI and URL values.
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_inline(p, text)
    return p


def shade_paragraph(paragraph, fill="F3F3F3"):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "12")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), "666666")
    p_bdr.append(left)
    ppr.append(p_bdr)


def add_caption(doc, text, kind):
    style = "Table Caption" if kind == "table" else "Figure Caption"
    p = doc.add_paragraph(style=style)
    add_inline(p, text, base_size=BODY_SIZE)
    return p


def strip_inline(text):
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    text = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", text)
    return text.strip()


def parse_md_table(lines, start):
    raw_rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        is_separator = all(re.fullmatch(r":?-{3,}:?", strip_inline(c).replace(" ", "")) for c in cells)
        if not is_separator:
            raw_rows.append(cells)
        i += 1
    if not raw_rows:
        return [], i
    cols = len(raw_rows[0])
    rows = [r[:cols] + [""] * max(0, cols - len(r)) for r in raw_rows]
    return rows, i


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


COMPACT_TABLE_HEADERS = {"stt", "mã", "id", "tt", "mức", "ưu tiên"}
STATUS_TABLE_TERMS = (
    "trạng thái",
    "kết quả",
    "mức hiện thực",
    "mức hiện tại",
    "mức ưu tiên",
)
NARRATIVE_TABLE_TERMS = (
    "mô tả",
    "nội dung",
    "ghi chú",
    "bằng chứng",
    "yêu cầu",
    "tiêu chí",
    "trách nhiệm",
    "điều kiện",
    "hiện vật",
    "artifact",
    "protocol",
    "vai trò",
    "lý do",
    "giới hạn",
    "bất nhất",
    "thiếu",
    "ảnh hưởng",
    "xử lý",
    "quy tắc",
)


def table_header_kind(header):
    normalized = re.sub(r"\s+", " ", str(header).casefold()).strip()
    if normalized in COMPACT_TABLE_HEADERS:
        return "compact"
    if any(term in normalized for term in STATUS_TABLE_TERMS):
        return "status"
    if any(term in normalized for term in NARRATIVE_TABLE_TERMS):
        return "narrative"
    return "general"


def compute_widths(rows, total=TABLE_DXA):
    """Allocate fixed table widths without collapsing labels/status columns.

    Character-count-only sizing made identifiers such as ``Cross-encoder`` and
    statuses wrap letter by letter.  Reserve a semantic minimum first, then
    distribute the remaining space by capped content length.
    """

    cols = len(rows[0])
    base_general = 1750 if cols <= 3 else 1350 if cols == 4 else 1100
    base_narrative = 2300 if cols <= 3 else 1650 if cols == 4 else 1350
    base_status = 1900 if cols <= 3 else 1550 if cols == 4 else 1250
    base_compact = 850 if cols <= 4 else 720
    token_cap = 2050 if cols <= 3 else 1800 if cols == 4 else 1450

    minima = []
    weights = []
    for idx in range(cols):
        values = [str(row[idx]) for row in rows]
        header = values[0]
        header_normalized = header.casefold().strip()
        kind = table_header_kind(header)
        max_len = max(4, min(64, max(len(value) for value in values)))

        if kind == "compact":
            minimum = base_compact
            weight = 6
        elif kind == "status":
            minimum = base_status
            weight = max(10, min(24, max_len))
        elif kind == "narrative":
            minimum = base_narrative
            weight = max(18, min(32, max_len))
        else:
            minimum = base_general
            weight = max(10, min(28, max_len))

        if "viết tắt" in header_normalized:
            minimum = max(minimum, 1550)
            weight = min(weight, 8)
        if "tầng" == header_normalized:
            minimum = max(minimum, 1500)
        if idx == 0 and kind != "compact":
            minimum = max(minimum, 1600 if cols <= 3 else 1350 if cols == 4 else 1200)

        # Protect ordinary labels/identifiers.  Narrative cells may contain
        # paths or URLs that must still wrap, so they do not drive this floor.
        if kind != "narrative":
            tokens = re.findall(r"[\w./:@-]+", " ".join(values), flags=re.UNICODE)
            longest_token = min(18, max((len(token) for token in tokens), default=0))
            minimum = max(minimum, min(token_cap, 250 + 100 * longest_token))

        minima.append(minimum)
        weights.append(weight)

    floor = 650
    while sum(minima) > total:
        candidates = [idx for idx, width in enumerate(minima) if width > floor]
        if not candidates:
            break
        excess = sum(minima) - total
        cut = max(1, (excess + len(candidates) - 1) // len(candidates))
        for idx in candidates:
            minima[idx] = max(floor, minima[idx] - cut)

    remaining = total - sum(minima)
    widths = list(minima)
    if remaining > 0:
        weight_sum = sum(weights)
        additions = [int(remaining * weight / weight_sum) for weight in weights]
        widths = [width + addition for width, addition in zip(widths, additions)]
    widths[-1] += total - sum(widths)
    return widths


def set_table_geometry(table, widths):
    table.autofit = False
    # Use an explicit left alignment plus a small indent.  This places the
    # near-full-width table at the visual center of the text block while also
    # surviving a Microsoft Word field-update/save cycle; Word drops tblInd
    # when the table itself is marked as centered.
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout"):
        remove_element(tbl_pr, tag)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            remove_element(tc_pr, "w:tcW")
            tc_w = OxmlElement("w:tcW")
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            tc_pr.append(tc_w)


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    widths = compute_widths(rows)
    set_table_geometry(table, widths)
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx]
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "true")
        tr_pr.append(cant_split)
        if r_idx == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
        for c_idx, value in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = r_idx == 0
            header_kind = table_header_kind(rows[0][c_idx])
            narrative = header_kind == "narrative" or len(value) > 20 or c_idx == len(row_data) - 1
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 or not narrative else WD_ALIGN_PARAGRAPH.LEFT
            plain_value = strip_inline(value)
            if (
                r_idx > 0
                and re.fullmatch(r"[\w./-]{1,12}", plain_value, flags=re.UNICODE)
                and 250 + 100 * len(plain_value) <= widths[c_idx]
            ):
                # Keep short identifiers such as XLSX, config and deployment
                # intact when a fixed-width table happens to cross a page.
                tc_pr = cell._tc.get_or_add_tcPr()
                remove_element(tc_pr, "w:noWrap")
                tc_pr.append(OxmlElement("w:noWrap"))
            # Prefer a semantic break at path/compound separators instead of
            # letting Word split the following identifier mid-word.
            display_value = re.sub(r"(?<=\w)/(?=\w)", "/\u200b", value)
            add_inline(p, display_value, base_size=TABLE_SIZE)
            for run in p.runs:
                set_run_font(run, size=TABLE_SIZE, bold=True if r_idx == 0 else None)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    spacer.paragraph_format.first_line_indent = Cm(0)
    return table


def add_image(doc, md_path, target, alt):
    target = unquote(target.strip())
    path = (md_path.parent / target).resolve()
    if not path.exists():
        raise FileNotFoundError(f"Không tìm thấy hình: {target}")
    max_width_cm = 15.5
    max_height_cm = 18.5
    width_cm = max_width_cm
    if Image is not None:
        with Image.open(path) as image:
            px_w, px_h = image.size
        if px_w and px_h:
            height_cm = width_cm * px_h / px_w
            if height_cm > max_height_cm:
                width_cm = max_height_cm * px_w / px_h
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(path), width=Cm(width_cm))
    try:
        doc_pr = shape._inline.docPr
        doc_pr.set("descr", alt)
        doc_pr.set("title", alt)
    except Exception:
        pass
    return p


def add_code_block(doc, code_lines):
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.keep_together = False
    text = "\n".join(code_lines)
    run = p.add_run(text)
    set_run_font(run, name=FONT_CODE, size=BODY_SIZE)
    shade_paragraph(p, "F4F4F4")
    return p


def add_toc_field(doc, kind):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.5
    if kind == "toc":
        instruction = ' TOC \\o "1-3" \\h \\z \\u \\t "Part Title,1" '
        display = "Mục lục sẽ được cập nhật tự động khi mở tài liệu trong Microsoft Word."
    elif kind == "tables":
        instruction = ' TOC \\h \\z \\t "Table Caption,1" '
        display = "Danh mục bảng sẽ được cập nhật tự động khi mở tài liệu trong Microsoft Word."
    else:
        instruction = ' TOC \\h \\z \\t "Figure Caption,1" '
        display = "Danh mục hình sẽ được cập nhật tự động khi mở tài liệu trong Microsoft Word."
    add_field(p, instruction, display)
    return p


def heading_style_for(text, level):
    upper = text.upper()
    if level == 1 and upper in FRONT_TITLES:
        return "Front Matter Title"
    return f"Heading {min(level, 4)}"


def build_report(md_path, docx_path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    cover = extract_cover_lines(lines)
    doc = Document()
    configure_styles(doc)
    set_update_fields(doc)
    add_cover(doc, cover)

    doc.core_properties.title = cover[2].title()
    doc.core_properties.subject = "Đồ án tốt nghiệp thạc sĩ - Hệ thống Thông tin"
    doc.core_properties.author = cover[1].title()
    doc.core_properties.keywords = "Enterprise AI Platform, LLM, RAG, AI Agent, MCP"

    i = 0
    in_cover = False
    in_signature = False
    have_content = False
    main_section_started = False
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()

        if line == "<!-- COVER -->":
            in_cover = True
            i += 1
            continue
        if line == "<!-- END_COVER -->":
            in_cover = False
            i += 1
            continue
        if line == "<!-- SIGNATURE -->":
            in_signature = True
            i += 1
            continue
        if line == "<!-- END_SIGNATURE -->":
            in_signature = False
            i += 1
            continue
        if in_cover:
            i += 1
            continue
        if in_signature:
            if line:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.keep_together = True
                add_inline(p, line)
            i += 1
            continue
        if line.startswith("<!--") or line == "" or re.fullmatch(r"-{3,}", line):
            i += 1
            continue

        if line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip())
                i += 1
            add_code_block(doc, code)
            have_content = True
            i += 1
            continue

        if line.startswith("[[EQUATION]]"):
            p = doc.add_paragraph(style="Equation")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            add_inline(p, line[len("[[EQUATION]]"):].strip())
            have_content = True
            i += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            text = strip_inline(heading.group(2))
            style_name = heading_style_for(text, level)
            page_break_before = False
            if level == 1:
                if text.upper() == MAIN_START_TITLE and not main_section_started:
                    main_section = doc.add_section(WD_SECTION.NEW_PAGE)
                    configure_section(main_section, cover=False)
                    set_page_border(main_section, False)
                    add_page_number(main_section, 1)
                    main_section_started = True
                elif have_content:
                    # A paragraph-level page break is ignored when the heading
                    # already lands at the top of a page.  A standalone break
                    # paragraph can otherwise create an empty page at an exact
                    # pagination boundary.
                    page_break_before = True
                if text.lower().startswith("chương 1"):
                    p_part = doc.add_paragraph("NỘI DUNG", style="Part Title")
                    # The following Heading 1 normally starts a new page.  Put
                    # that break on the part label instead so NỘI DUNG cannot
                    # be orphaned at the bottom of the preceding MỞ ĐẦU page.
                    p_part.paragraph_format.page_break_before = True
                    p_part.paragraph_format.keep_with_next = True
                    page_break_before = False
                    for run in p_part.runs:
                        set_run_font(run, size=BODY_SIZE, bold=True)
            p = doc.add_paragraph(style=style_name)
            p.paragraph_format.first_line_indent = Cm(0)
            if page_break_before:
                p.paragraph_format.page_break_before = True
            add_inline(p, text, base_size=BODY_SIZE)
            for run in p.runs:
                set_run_font(run, size=BODY_SIZE, bold=True, italic=True if level == 3 else None)
            have_content = True
            i += 1
            continue

        if line == "[[TOC]]":
            add_toc_field(doc, "toc")
            have_content = True
            i += 1
            continue
        if line == "[[LIST_OF_TABLES]]":
            add_toc_field(doc, "tables")
            have_content = True
            i += 1
            continue
        if line == "[[LIST_OF_FIGURES]]":
            add_toc_field(doc, "figures")
            have_content = True
            i += 1
            continue

        if line.startswith("|") and line.endswith("|"):
            rows, next_i = parse_md_table(lines, i)
            caption = None
            j = next_i
            while j < len(lines) and not lines[j].strip():
                j += 1
            if j < len(lines):
                cap_match = re.fullmatch(r"\*(Bảng\s+\d+(?:\.\d+)*[.:]\s*.+)\*", lines[j].strip(), re.IGNORECASE)
                if cap_match:
                    caption = cap_match.group(1)
                    next_i = j + 1
            if caption:
                add_caption(doc, caption, "table")
            if rows:
                add_table(doc, rows)
            have_content = True
            i = next_i
            continue

        image = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", line)
        if image:
            add_image(doc, md_path, image.group(2), image.group(1))
            have_content = True
            i += 1
            continue

        cap = re.fullmatch(r"\*(Hình|Bảng)\s+(\d+(?:\.\d+)*)([.:])\s*(.+)\*", line, re.IGNORECASE)
        if cap:
            text = f"{cap.group(1).capitalize()} {cap.group(2)}{cap.group(3)} {cap.group(4)}"
            add_caption(doc, text, "figure" if cap.group(1).lower() == "hình" else "table")
            have_content = True
            i += 1
            continue

        bullet = re.match(r"^[-+]\s+(.+)$", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.first_line_indent = None
            add_inline(p, bullet.group(1))
            have_content = True
            i += 1
            continue

        numbered = re.match(r"^(\d+)[.)]\s+(.+)$", line)
        if numbered:
            # Preserve the explicit number from Markdown.  Word's built-in
            # List Number style may restart after a page boundary/field update,
            # which previously turned a source 1–4 sequence into 1, 1, 2.
            p = doc.add_paragraph(style="Manual Numbered List")
            prefix = p.add_run(f"{numbered.group(1)}. ")
            set_run_font(prefix, size=BODY_SIZE)
            add_inline(p, numbered.group(2))
            have_content = True
            i += 1
            continue

        if line.startswith(">"):
            p = doc.add_paragraph(style="Note")
            add_inline(p, line.lstrip("> "))
            shade_paragraph(p, "F4F4F4")
            have_content = True
            i += 1
            continue

        if re.match(r"^\[\d+\]", line):
            add_body_paragraph(doc, line, style="Reference", indent=False)
        else:
            add_body_paragraph(doc, line)
        have_content = True
        i += 1

    if not main_section_started:
        raise RuntimeError(f"Không tìm thấy tiêu đề bắt đầu đánh số trang: {MAIN_START_TITLE}")

    OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)
    print(f"[OK] Da tao: {docx_path}")


if __name__ == "__main__":
    if not INPUT_FILE.exists():
        raise FileNotFoundError(f"Chưa có bản Markdown đã gộp: {INPUT_FILE}")
    chapter_files = list((REPORT_DIR / "chapters").glob("*.md"))
    if chapter_files and INPUT_FILE.stat().st_mtime_ns < max(p.stat().st_mtime_ns for p in chapter_files):
        raise RuntimeError("Bản Markdown đã gộp cũ hơn file chương; hãy chạy merge_report.py trước")
    build_report(INPUT_FILE, OUTPUT_FILE)
