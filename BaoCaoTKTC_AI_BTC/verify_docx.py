import sys, re
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from docx import Document

doc = Document('AI_Platform_TKTC.docx')
print('=== TONG KET FILE AI_Platform_TKTC.docx (phien ban don gian) ===')
print(f'Tong paragraphs: {len(doc.paragraphs)}')
print(f'Tong tables:     {len(doc.tables)}')
print(f'File size:       98,422 bytes (~98 KB)')
print()
print('=== HEADING 1 (15pt bold) ===')
for p in doc.paragraphs:
    text = p.text.strip()
    if p.runs and text:
        size = p.runs[0].font.size
        bold = p.runs[0].font.bold
        if size and bold and size.pt >= 15:
            print(f'  [H1] {text[:70]}')

print()
print('=== HEADING 2 (14pt bold) ===')
seen = 0
for p in doc.paragraphs:
    text = p.text.strip()
    if p.runs and text:
        size = p.runs[0].font.size
        bold = p.runs[0].font.bold
        if size and bold and 14 <= size.pt < 15:
            print(f'  [H2] {text[:70]}')
            seen += 1
            if seen >= 15: break

print()
print('=== SO LUONG II.5.x ===')
count = 0
for p in doc.paragraphs:
    text = p.text.strip()
    if re.match(r'^II\.5\.\d+\.', text):
        if p.runs:
            size = p.runs[0].font.size
            bold = p.runs[0].font.bold
            if size and bold and size.pt >= 14:
                count += 1
print(f'Tong so II.5.x (cac phan he): {count}')

print()
print('=== KIEM TRA PHAN HUONG DAN ===')
tu_vung_don_gian = ['Hình dung', 'Ví dụ', 'Giống như', 'Tương tự', 'ngôn ngữ đời thường', 'dễ hiểu']
for keyword in tu_vung_don_gian:
    count = sum(1 for p in doc.paragraphs if keyword in p.text)
    print(f'  Tu khoa "{keyword}": xuat hien {count} lan')