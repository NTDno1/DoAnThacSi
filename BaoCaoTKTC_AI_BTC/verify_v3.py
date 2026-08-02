import sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from docx import Document

doc = Document('AI_Platform_TKTC_v3.docx')
print('=== TONG KET FILE AI_Platform_TKTC_v3.docx (v3) ===')
print(f'Tong paragraphs: {len(doc.paragraphs)}')
print(f'Tong tables:     {len(doc.tables)}')
print(f'File size:       117,408 bytes')
print()

# So do ASCII art
count = sum(1 for p in doc.paragraphs if any(c in p.text for c in ['╔', '┌', '★', '▼', '►']))
print(f'Paragraphs co ASCII art: {count}')

# Cac heading II.4.8 x moi them
print()
print('=== HEADING II.4.8 (CAC MUC MOI THEM) ===')
import re
for p in doc.paragraphs:
    text = p.text.strip()
    if re.match(r'^II\.4\.8\.\d+', text):
        if p.runs:
            size = p.runs[0].font.size
            bold = p.runs[0].font.bold
            if size and bold:
                print(f'  [size={size.pt}] {text[:80]}')

# Cac bang co nhieu dong (tables dung cho thong so)
print()
print('=== CAC BANG LON TRONG FILE (so dong) ===')
big_tables = sorted([(len(t.rows), i+1) for i, t in enumerate(doc.tables)], reverse=True)[:10]
for n_rows, idx in big_tables:
    print(f'  Table #{idx}: {n_rows} rows')